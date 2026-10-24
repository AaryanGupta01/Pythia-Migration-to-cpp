"""Tiny python-docx DSL for the final report: headings, paragraphs with **bold**/`code`/*italic* markup, tables with
captions, figures with captions, code blocks, landscape sections, TOC and page-number fields.
Cross-references: write [[fig:key]] / [[tab:key]] in text; numbers are assigned in a first pass (N.k = chapter.counter)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ACCENT = RGBColor(0x3B, 0x30, 0xC4)
INK = RGBColor(0x14, 0x17, 0x1F)
GREY = RGBColor(0x5B, 0x62, 0x70)
BODY_FONT, MONO_FONT = "Calibri", "Consolas"
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)")


def _shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def _cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def _borders(table, color="C9CCD6", sz=4, inside=True):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for k in edges:
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _field(par, instr, placeholder=""):
    r = par.add_run()
    for t, txt in (("begin", None), (None, instr), ("separate", None), (None, placeholder), ("end", None)):
        if t:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), t)
            r._r.append(fc)
        elif txt is instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            r._r.append(it)
        else:
            tt = OxmlElement("w:t")
            tt.text = placeholder
            r._r.append(tt)
    return r


class Report:
    def __init__(self, header_text: str, refs: dict | None = None):
        self.doc = Document()
        self.refs = refs or {}          # filled by the first pass
        self.new_refs: dict = {}
        self.h1n = 0
        self.fig_n = self.tab_n = 0
        self.app_letter = None
        self.header_text = header_text
        self._setup()

    # ------------------------------------------------------------------ setup
    def _setup(self):
        d = self.doc
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.9)
        st = d.styles
        n = st["Normal"]
        n.font.name = BODY_FONT
        n.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        n.font.size = Pt(10)
        n.paragraph_format.space_after = Pt(4)
        n.paragraph_format.line_spacing = 1.07
        for name, size, before, after, color in (("Heading 1", 17, 16, 6, ACCENT), ("Heading 2", 12.5, 11, 3, INK),
                                                 ("Heading 3", 10.8, 8, 2, GREY)):
            h = st[name]
            h.font.name = BODY_FONT
            h.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            h.element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
            h.element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = color
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            st[name].font.name = BODY_FONT
            st[name].paragraph_format.space_after = Pt(2)
        self._header_footer(s)

    def _header_footer(self, s, first_blank=True):
        s.different_first_page_header_footer = first_blank
        hp = s.header.paragraphs[0]
        hp.text = ""
        r = hp.add_run(self.header_text)
        r.font.size, r.font.color.rgb = Pt(8.5), GREY
        pPr = hp._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        bt = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "C9CCD6")):
            bt.set(qn(f"w:{k}"), v)
        bd.append(bt)
        pPr.append(bd)
        fp = s.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Page ")
        r.font.size, r.font.color.rgb = Pt(8.5), GREY
        pr = _field(fp, "PAGE", "1")
        pr.font.size, pr.font.color.rgb = Pt(8.5), GREY

    # ------------------------------------------------------------------ text
    def _runs(self, par, text, size=None, color=None, bold=None, italic=None, mono_size=None):
        text = self._resolve(text)
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = par.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = MONO_FONT
                r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
                r.font.size = Pt(mono_size or ((size or 10.5) - 1))
                r.font.color.rgb = RGBColor(0x2E, 0x2A, 0x9E)
                continue
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = par.add_run(part[1:-1])
                r.italic = True
            else:
                r = par.add_run(part)
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
            if bold:
                r.bold = True
            if italic:
                r.italic = True

    def _resolve(self, text):
        def sub(m):
            kind, key = m.group(1), m.group(2)
            ref = self.refs.get(f"{kind}:{key}")
            if ref is None:
                return {"fig": "Figure ?", "tab": "Table ?"}[kind]
            return ("Figure " if kind == "fig" else "Table ") + ref
        return re.sub(r"\[\[(fig|tab):([\w-]+)\]\]", sub, text)

    def p(self, text, align=None, size=None, after=None, color=None, italic=None, keep=False):
        par = self.doc.add_paragraph()
        self._runs(par, text, size=size, color=color, italic=italic)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        if keep:
            par.paragraph_format.keep_with_next = True
        return par

    def bullets(self, items, style="List Bullet", size=None):
        for it in items:
            par = self.doc.add_paragraph(style=style)
            self._runs(par, it, size=size)
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.45)

    def numbered(self, items, start=1):
        for i, it in enumerate(items, start):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            par.paragraph_format.space_after = Pt(2)
            self._runs(par, f"**{i}.** " + it)

    def h1(self, text, numbered=True, page_break=False):
        if page_break:
            self.pagebreak()
        if numbered:
            self.h1n += 1
            self.fig_n = self.tab_n = 0
            label = f"{self.h1n}. {text}"
        else:
            label = text
        self.cur_prefix = str(self.h1n) if numbered else (self.app_letter or "A")
        return self.doc.add_heading(label, level=1)

    def appendix(self, letter, title):
        self.app_letter = letter
        self.fig_n = self.tab_n = 0
        self.cur_prefix = letter
        self.pagebreak()
        return self.doc.add_heading(f"Appendix {letter}. {title}", level=1)

    def h2(self, text):
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    def pagebreak(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ captions
    def _caption(self, kind, key, text):
        if kind == "fig":
            self.fig_n += 1
            num = f"{self.cur_prefix}.{self.fig_n}"
            label = "Figure"
        else:
            self.tab_n += 1
            num = f"{self.cur_prefix}.{self.tab_n}"
            label = "Table"
        self.new_refs[f"{kind}:{key}"] = num
        par = self.doc.add_paragraph()
        par.paragraph_format.space_after = Pt(8) if kind == "fig" else Pt(3)
        par.paragraph_format.space_before = Pt(2) if kind == "fig" else Pt(6)
        par.paragraph_format.keep_with_next = kind == "tab"
        if kind == "fig":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(f"{label} {num}. ")
        r.bold = True
        r.font.size, r.font.color.rgb = Pt(9), ACCENT
        self._runs(par, text, size=9, color=GREY)

    def figure(self, path, caption, key, width_cm=16.0):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(str(path), width=Cm(width_cm))
        self._caption("fig", key, caption)

    def table(self, header, rows, key, caption, widths=None, size=8.8, align=None, first_col_bold=False):
        if key:
            self._caption("tab", key, caption)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _borders(t)
        _cell_margins(t)
        total = 17.0
        widths = widths or [total / len(header)] * len(header)
        scale = total / sum(widths)
        widths = [w * scale for w in widths]
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.width = Cm(widths[i])
            _shade(c, "ECEBFB")
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            self._runs(c.paragraphs[0], h, size=size, bold=True, color=INK)
        tr = t.rows[0]._tr
        trPr = tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for row in rows:
            cells = t.add_row().cells
            rp = t.rows[-1]._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            rp.append(cs)
            for i, v in enumerate(row):
                c = cells[i]
                c.width = Cm(widths[i])
                par = c.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                lines = str(v).split("\n")
                for li, ln in enumerate(lines):
                    if li:
                        par = c.add_paragraph()
                        par.paragraph_format.space_after = Pt(0)
                        par.paragraph_format.line_spacing = 1.0
                    self._runs(par, ln, size=size, bold=(first_col_bold and i == 0))
                    if align and align[i] == "r":
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    elif align and align[i] == "c":
                        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 0.6
        return t

    def table_plain(self, header, rows, widths, size=9.5):
        return self.table(header, rows, None, None, widths=widths, size=size)

    def code(self, text, size=8, caption=None, key=None):
        if caption:
            self._caption("tab", key or caption[:8], caption) if False else None
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, "D6D9E3", 4, inside=False)
        _cell_margins(t, 60, 60, 110, 110)
        c = t.rows[0].cells[0]
        c.width = Cm(16.6)
        _shade(c, "F6F7FB")
        lines = text.strip("\n").split("\n")
        for i, ln in enumerate(lines):
            par = c.paragraphs[0] if i == 0 else c.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            r = par.add_run(ln if ln else " ")
            r.font.name = MONO_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0x1F, 0x23, 0x33)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)
        sp.paragraph_format.line_spacing = 0.6

    def callout(self, text, color="ECEBFB"):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, "B9B4F5", 6, inside=False)
        _cell_margins(t, 70, 70, 130, 130)
        c = t.rows[0].cells[0]
        c.width = Cm(16.6)
        _shade(c, color)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        self._runs(c.paragraphs[0], text, size=10)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)
        sp.paragraph_format.line_spacing = 0.6

    def toc(self, levels="1-2"):
        par = self.doc.add_paragraph()
        _field(par, f'TOC \\o "{levels}" \\h \\z \\u', "Right-click and update field to build the table of contents.")

    def landscape_section(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.6)
        s.top_margin, s.bottom_margin = Cm(1.8), Cm(1.6)
        s.different_first_page_header_footer = False
        return s

    def portrait_section(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin, s.bottom_margin = Cm(2.2), Cm(2.0)
        s.different_first_page_header_footer = False
        return s

    def save(self, path):
        self.doc.save(str(path))
